"""The applications of a workshop as an Excel file, one column per question (PRD §7.6)."""

import io
from typing import Any

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from workshop_manager.exports.reports import answer_for, applications

HEADER_FILL = PatternFill("solid", fgColor="283878")
BASE_COLUMNS = [
    ("Nazwisko", 18),
    ("Imię", 14),
    ("E-mail", 28),
    ("Telefon", 15),
    ("Poziom", 16),
    ("Status", 16),
    ("Miejsce na liście rezerwowej", 12),
    ("Data zgłoszenia", 17),
]
TAIL_COLUMNS = [
    ("Uwagi uczestnika", 30),
    ("Notatki organizatora", 30),
    ("Zgoda na informacje o warsztatach", 12),
]

#: Cells starting with these could run as formulas when the file is opened (CSV/formula
#: injection); a leading apostrophe keeps them text.
FORMULA_START = ("=", "+", "-", "@")


def _text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


def applications_xlsx(workshop: Any, statuses: list[str] | None = None) -> bytes:
    questions = list(workshop.questions.order_by("order", "pk"))
    book = Workbook()
    sheet = book.active
    sheet.title = "Zgłoszenia"
    headers = [*BASE_COLUMNS, *((q.label, 30) for q in questions), *TAIL_COLUMNS]
    sheet.append([title for title, _width in headers])
    for index, (_title, width) in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for application in applications(workshop, statuses):
        submitted = timezone.localtime(application.submitted_at).replace(tzinfo=None)
        row = [
            application.last_name,
            application.first_name,
            application.email,
            application.phone,
            application.level.name,
            application.get_status_display(),
            application.waitlist_position,
            submitted,
            *(answer_for(application, question) for question in questions),
            application.remarks,
            application.admin_notes,
            "tak" if application.marketing_consent else "nie",
        ]
        sheet.append([_text(value) for value in row])
        sheet.cell(row=sheet.max_row, column=8).number_format = "yyyy-mm-dd hh:mm"

    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
