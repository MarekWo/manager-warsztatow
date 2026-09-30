"""Messages to participants, the Excel export and the printouts (PRD §7.5, §7.6)."""

import io
from datetime import date, time

import pytest
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from tests.factories import (
    AdminFactory,
    ApplicationFactory,
    LevelFactory,
    QuestionFactory,
    SessionFactory,
    WorkshopFactory,
)
from workshop_manager.applications.models import Status
from workshop_manager.communications import services
from workshop_manager.communications.models import Broadcast, EmailMessage, Group
from workshop_manager.exports import pdf, reports
from workshop_manager.forms_builder.models import QuestionKind


@pytest.fixture
def admin_client(client, db):
    client.force_login(AdminFactory(email="admin@example.com"))
    return client


@pytest.fixture
def workshop(db):
    workshop = WorkshopFactory(title="Rekolekcje z ikoną")
    SessionFactory(
        workshop=workshop, date=date(2026, 11, 28), start_time=time(10), end_time=time(16)
    )
    SessionFactory(
        workshop=workshop, date=date(2026, 11, 29), start_time=time(9), end_time=time(13)
    )
    LevelFactory(workshop=workshop, name="Początkujący", order=0)
    LevelFactory(workshop=workshop, name="Zaawansowani", order=1)
    return workshop


def level(workshop, name):
    return workshop.levels.get(name=name)


# --- Messages -------------------------------------------------------------------------------------


def test_message_to_the_accepted_with_an_attachment(admin_client, workshop):
    beginners = level(workshop, "Początkujący")
    anna = ApplicationFactory(level=beginners, first_name="Anna", status=Status.ACCEPTED)
    ApplicationFactory(level=beginners, status=Status.WAITLISTED)
    ApplicationFactory(level=level(workshop, "Zaawansowani"), status=Status.ACCEPTED)

    response = admin_client.post(
        reverse("panel:broadcast_add", args=[workshop.pk]),
        {
            "group": Group.ACCEPTED,
            "subject": "Co zabrać: {warsztat}",
            "body": "Dzień dobry {imie},\n\nw załączniku lista materiałów.",
            "attachment": SimpleUploadedFile(
                "materiały.pdf", b"%PDF-1.4 test", content_type="application/pdf"
            ),
        },
    )
    broadcast = Broadcast.objects.get()
    preview_url = reverse("panel:broadcast_preview", args=[workshop.pk, broadcast.pk])
    assert response["Location"] == preview_url
    assert not EmailMessage.objects.exists()  # nothing goes out before the preview

    page = admin_client.get(preview_url).content.decode()
    assert "Wyślij do 2 os." in page
    assert "materiały.pdf" in page

    admin_client.post(reverse("panel:broadcast_send", args=[workshop.pk, broadcast.pk]))
    emails = EmailMessage.objects.filter(broadcast=broadcast)
    assert emails.count() == 2
    personal = emails.get(to_email=anna.email)
    assert personal.subject == "Co zabrać: Rekolekcje z ikoną"
    assert "Dzień dobry Anna" in personal.body_text

    # Sending again changes nothing.
    admin_client.post(reverse("panel:broadcast_send", args=[workshop.pk, broadcast.pk]))
    assert EmailMessage.objects.count() == 2

    services.send_email_message(personal.pk)
    sent = mail.outbox[-1]
    assert sent.attachments[0][0] == "materiały.pdf"
    assert sent.attachments[0][1] == b"%PDF-1.4 test"


def test_one_email_per_address_and_level_group(workshop, db):
    beginners = level(workshop, "Początkujący")
    ApplicationFactory(level=beginners, status=Status.ACCEPTED)
    ApplicationFactory(level=level(workshop, "Zaawansowani"), status=Status.ACCEPTED)
    broadcast = Broadcast.objects.create(
        workshop=workshop, group=Group.LEVEL, level=beginners, subject="S", body="B"
    )
    from workshop_manager.communications.broadcasts import recipients

    assert [a.level for a in recipients(broadcast)] == [beginners]


def test_message_form_checks_placeholders_and_level(admin_client, workshop):
    response = admin_client.post(
        reverse("panel:broadcast_add", args=[workshop.pk]),
        {"group": Group.LEVEL, "subject": "Temat", "body": "Dzień dobry {imię}"},
    )
    page = response.content.decode()
    assert "Nieznane pola: {imię}" not in page  # braces with Polish letters are plain text
    assert "Wybierz poziom." in page
    response = admin_client.post(
        reverse("panel:broadcast_add", args=[workshop.pk]),
        {"group": Group.ACCEPTED, "subject": "Temat", "body": "Dzień dobry {imi}"},
    )
    assert "Nieznane pola: {imi}" in response.content.decode()
    assert not Broadcast.objects.exists()


def test_message_to_selected_people_from_the_list(admin_client, workshop):
    beginners = level(workshop, "Początkujący")
    chosen = ApplicationFactory(level=beginners)
    ApplicationFactory(level=beginners)
    response = admin_client.post(
        reverse("panel:application_bulk"), {"ids": [chosen.pk], "action": "message"}
    )
    broadcast = Broadcast.objects.get()
    assert broadcast.group == Group.SELECTED
    assert response["Location"] == reverse("panel:broadcast_edit", args=[workshop.pk, broadcast.pk])
    admin_client.post(
        response["Location"], {"group": Group.SELECTED, "subject": "Zmiana sali", "body": "Treść"}
    )
    admin_client.post(reverse("panel:broadcast_send", args=[workshop.pk, broadcast.pk]))
    assert list(EmailMessage.objects.values_list("to_email", flat=True)) == [chosen.email]


# --- Exports and printouts ----------------------------------------------------------------------


def test_excel_export_has_a_column_per_question(admin_client, workshop):
    question = QuestionFactory(workshop=workshop, label="Doświadczenie", order=1)
    accepted = ApplicationFactory(
        level=level(workshop, "Początkujący"), status=Status.ACCEPTED, last_name="Żółć"
    )
    accepted.answers.create(question=question, label=question.label, value="=HYPERLINK(1)")
    ApplicationFactory(level=level(workshop, "Początkujący"), status=Status.REJECTED)

    response = admin_client.get(reverse("panel:export_xlsx", args=[workshop.pk]))
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    sheet = load_workbook(io.BytesIO(response.content)).active
    header = [cell.value for cell in sheet[1]]
    assert "Doświadczenie" in header
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1  # accepted only by default
    assert rows[0][0] == "Żółć"
    assert rows[0][header.index("Doświadczenie")] == "'=HYPERLINK(1)"

    everyone = admin_client.get(reverse("panel:export_xlsx", args=[workshop.pk]), {"zakres": "all"})
    assert load_workbook(io.BytesIO(everyone.content)).active.max_row == 3


def test_attendance_list_has_the_accepted_and_a_column_per_session(admin_client, workshop):
    ApplicationFactory(
        level=level(workshop, "Początkujący"), status=Status.ACCEPTED, last_name="Obecna"
    )
    ApplicationFactory(
        level=level(workshop, "Początkujący"), status=Status.REJECTED, last_name="Nieobecna"
    )
    page = admin_client.get(reverse("panel:attendance", args=[workshop.pk])).content.decode()
    assert "Obecna" in page
    assert "Nieobecna" not in page
    assert page.count('class="sign"') >= 2
    assert "28.11" in page
    assert "29.11" in page


@pytest.mark.skipif(not pdf.available(), reason="WeasyPrint needs Pango (present in Docker)")
def test_attendance_pdf(admin_client, workshop):
    ApplicationFactory(
        level=level(workshop, "Początkujący"), status=Status.ACCEPTED, last_name="Żółć"
    )
    response = admin_client.get(reverse("panel:attendance", args=[workshop.pk]), {"format": "pdf"})
    assert response["Content-Type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


def test_materials_summary_and_filter(admin_client, workshop):
    beginners = level(workshop, "Początkujący")
    board = QuestionFactory(
        workshop=workshop,
        label="Deska 26×26",
        kind=QuestionKind.MATERIAL,
        choices="Tak, zamawiam\nMam własną",
        level=beginners,
    )
    for status, value in [
        (Status.ACCEPTED, "Tak, zamawiam"),
        (Status.ACCEPTED, "Tak, zamawiam"),
        (Status.ACCEPTED, "Mam własną"),
        (Status.WAITLISTED, "Tak, zamawiam"),
        (Status.REJECTED, "Tak, zamawiam"),
    ]:
        application = ApplicationFactory(level=beginners, status=status)
        application.answers.create(question=board, label=board.label, value=value)

    summary = reports.materials(workshop)[0]
    order, own = summary.lines
    assert (order.choice, order.accepted, order.waiting) == ("Tak, zamawiam", 2, 1)
    assert own.accepted == 1
    page = admin_client.get(reverse("panel:materials", args=[workshop.pk])).content.decode()
    assert "Deska 26×26" in page

    url = reverse("panel:workshop_applications", args=[workshop.pk])
    listed = admin_client.get(url, {"material": f"{board.pk}|Mam własną"}).content.decode()
    assert listed.count('name="ids"') == 1


def test_report_pages_open(admin_client, workshop):
    for name in ("report_index", "contacts", "broadcast_list", "broadcast_add"):
        assert admin_client.get(reverse(f"panel:{name}", args=[workshop.pk])).status_code == 200
